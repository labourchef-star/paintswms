import os, json, time, secrets, hashlib, hmac, html, io, smtplib, datetime
from email.message import EmailMessage
from contextlib import contextmanager
import httpx
from fastapi import FastAPI, Request, Form
from fastapi.responses import HTMLResponse, RedirectResponse, Response, StreamingResponse
from . import content as C
from .disposable import is_disposable

SECRET=os.environ.get("APP_SECRET") or secrets.token_hex(32)
GUMROAD_PRODUCTS=[p for p in os.environ.get("GUMROAD_PRODUCT_IDS","u2Tu8n-yyl4toTk914y6jA==,dH-NHPOh2YK74-gz4J8DwA==").split(",") if p]
PACK_ID=os.environ.get("GUMROAD_PACK_ID","u2Tu8n-yyl4toTk914y6jA==")
PRO_URL="https://labourchef.gumroad.com/l/gztwsf"; PACK_URL="https://labourchef.gumroad.com/l/lldel"; FREE_URL="https://labourchef.gumroad.com/l/hjpom"
TS_SITE=os.environ.get("TURNSTILE_SITE_KEY",""); TS_SECRET=os.environ.get("TURNSTILE_SECRET","")
DEV=os.environ.get("DEV_SHOW_CODES")=="1"
app=FastAPI()

from .dbx import db, SCHEMA
with db() as c: c.executescript(SCHEMA)

# ---------- Funnel counters: day+event counts only, no PII ----------
EVENTS=("visit_home","preview_generated","email_captured","checkout_click_pack","checkout_click_unlimited","unlock_success")
def track(ev):
    if ev not in EVENTS: return
    try:
        with db() as c: c.execute("insert into counters(day,event,n) values(?,?,1) on conflict(day,event) do update set n=counters.n+1",(datetime.datetime.utcnow().strftime("%Y-%m-%d"),ev))
    except Exception: pass
TRACK_JS="""<script>document.addEventListener('click',function(e){var a=e.target.closest&&e.target.closest('a[href*="gumroad.com/l/"]');if(!a)return;var ev=a.href.indexOf('/l/lldel')>-1?'checkout_click_pack':a.href.indexOf('/l/gztwsf')>-1?'checkout_click_unlimited':'';if(ev&&navigator.sendBeacon)navigator.sendBeacon('/t/'+ev);},true);</script>"""


def now(): return time.time()
def esc(s): return html.escape(str(s or ""))
def ip(req): return (req.headers.get("x-forwarded-for","").split(",")[0].strip() or (req.client.host if req.client else "?"))
def device(req): return req.cookies.get("did") or "none"
def limited(keys,kind,n,win):
    with db() as c:
        for k in keys:
            if c.execute("select count(*) as n from hits where key=? and kind=? and ts>?",(k,kind,now()-win)).fetchone()["n"]>=n: return True
        for k in keys: c.execute("insert into hits values(?,?,?)",(k,kind,now()))
    return False
def sign(d): s=json.dumps(d,separators=(",",":")); return s+"."+hmac.new(SECRET.encode(),s.encode(),"sha256").hexdigest()
def unsign(t):
    try:
        s,m=t.rsplit(".",1)
        return json.loads(s) if hmac.compare_digest(m,hmac.new(SECRET.encode(),s.encode(),"sha256").hexdigest()) else None
    except Exception: return None
def turnstile_ok(token,req):
    if not TS_SECRET: return True
    try: return httpx.post("https://challenges.cloudflare.com/turnstile/v0/siteverify",data={"secret":TS_SECRET,"response":token,"remoteip":ip(req)},timeout=8).json().get("success",False)
    except Exception: return False
BASE_URL=os.environ.get("BASE_URL","https://paintswms.kerbform.com").rstrip("/")
BASE_HOST=BASE_URL.split("://",1)[-1]
def email_ready(): return bool(os.environ.get("RESEND_API_KEY") or os.environ.get("SMTP_HOST"))
def send_code(to,code):
    body=f"Your PaintSWMS code is {code}. It expires in 15 minutes. If you didn't ask for this, ignore this email."
    if os.environ.get("RESEND_API_KEY"):
        r=httpx.post("https://api.resend.com/emails",headers={"Authorization":"Bearer "+os.environ["RESEND_API_KEY"]},json={"from":os.environ.get("MAIL_FROM","PaintSWMS <codes@send.kerbform.com>"),"to":[to],"subject":f"PaintSWMS code: {code}","text":body},timeout=10); return r.status_code<300
    if os.environ.get("SMTP_HOST"):
        m=EmailMessage(); m["From"]=os.environ.get("MAIL_FROM"); m["To"]=to; m["Subject"]=f"PaintSWMS code: {code}"; m.set_content(body)
        with smtplib.SMTP(os.environ["SMTP_HOST"],int(os.environ.get("SMTP_PORT","587"))) as s:
            s.starttls(); s.login(os.environ["SMTP_USER"],os.environ["SMTP_PASS"]); s.send_message(m)
        return True
    return DEV
def make_code(email,purpose,payload=""):
    code=f"{secrets.randbelow(10**6):06d}"
    with db() as c:
        c.execute("delete from codes where email=? and purpose=?",(email,purpose))
        c.execute("insert into codes(email,purpose,hash,exp,payload) values(?,?,?,?,?)",(email,purpose,hashlib.sha256((SECRET+code).encode()).hexdigest(),now()+900,payload))
    return code
def check_code(email,purpose,code):
    with db() as c:
        r=c.execute("select * from codes where email=? and purpose=?",(email,purpose)).fetchone()
        if not r or r["exp"]<now() or r["tries"]>=5: return None
        c.execute("update codes set tries=tries+1 where email=? and purpose=?",(email,purpose))
        if hmac.compare_digest(r["hash"],hashlib.sha256((SECRET+code.strip()).encode()).hexdigest()):
            c.execute("delete from codes where email=? and purpose=?",(email,purpose)); return r["payload"] or "{}"
    return None
def user(req):
    t=req.cookies.get("sess")
    if not t: return None
    with db() as c:
        r=c.execute("select u.* from sessions s join users u on u.id=s.user_id where s.tok=? and s.exp>?",(hashlib.sha256(t.encode()).hexdigest(),now())).fetchone()
    return r
def login_resp(email,to="/app"):
    with db() as c:
        c.execute("insert into users(email,created) values(?,?) on conflict do nothing",(email,now()))
        uid=c.execute("select id from users where email=?",(email,)).fetchone()["id"]
        t=secrets.token_urlsafe(32); c.execute("insert into sessions values(?,?,?)",(hashlib.sha256(t.encode()).hexdigest(),uid,now()+30*86400))
    r=RedirectResponse(to,303); r.set_cookie("sess",t,max_age=30*86400,httponly=True,samesite="lax",secure=os.environ.get("SECURE_COOKIES","1")=="1"); return r
def is_pro(u): return bool(u and u["pro_until"]>now())
def verify_license(key):
    for pid in GUMROAD_PRODUCTS:
        try: j=httpx.post("https://api.gumroad.com/v2/licenses/verify",data={"product_id":pid,"license_key":key.strip(),"increment_uses_count":"false"},timeout=10).json()
        except Exception: continue
        if j.get("success"):
            p=j.get("purchase",{})
            if p.get("refunded") or p.get("chargebacked") or p.get("subscription_cancelled_at") or p.get("subscription_failed_at"): return None
            return pid,p
    return None

CSS="""<meta name=viewport content="width=device-width,initial-scale=1"><style>
:root{--ink:#1d2430;--acc:#e8572a;--bg:#f7f5f1}*{box-sizing:border-box}body{margin:0;font:16px/1.5 system-ui,Segoe UI,Arial,sans-serif;color:var(--ink);background:var(--bg)}
.w{max-width:860px;margin:auto;padding:0 14px}header{background:var(--ink);color:#fff;padding:12px 0}header .w{display:flex;justify-content:space-between;align-items:center}header a{color:#fff;text-decoration:none}.logo{font-weight:800;font-size:20px}.logo span{color:var(--acc)}
.card{background:#fff;border:1px solid #e4e0d8;border-radius:10px;padding:18px;margin:16px 0}label{display:block;font-weight:700;margin:12px 0 5px}input,select,textarea{width:100%;padding:12px;border:1px solid #bbb;border-radius:8px;font-size:16px}
.seg{display:flex;gap:8px;flex-wrap:wrap}.seg label{flex:1;min-width:70px;margin:0;font-weight:600;border:1px solid #bbb;border-radius:8px;padding:11px;text-align:center;background:#fff}.seg input{display:none}.seg input:checked+span{color:var(--acc)}.seg label:has(input:checked){border-color:var(--acc);background:#fff4ef}
.btn{display:inline-block;background:var(--acc);color:#fff!important;padding:14px 20px;border-radius:8px;font-weight:700;text-decoration:none;border:0;cursor:pointer;font-size:17px;width:100%;text-align:center}.btn.sec{background:var(--ink)}.btn.sm{width:auto;padding:9px 14px;font-size:15px}
table{width:100%;border-collapse:collapse;font-size:13.5px;margin:8px 0}td,th{border:1px solid #999;padding:6px;vertical-align:top;text-align:left}th{background:#eee}
.wm{position:relative;overflow:hidden}.wm:after{content:"FREE PREVIEW — paintswms";position:absolute;top:40%;left:-10%;width:120%;text-align:center;transform:rotate(-25deg);font-size:44px;font-weight:800;color:rgba(232,87,42,.18);pointer-events:none}
.lock{filter:blur(4px);user-select:none}.hp{position:absolute;left:-5000px}.err{background:#fde8e8;border:1px solid #f3b4b4;padding:10px;border-radius:8px}.ok{background:#eaf6ea;border:1px solid #b6deb6;padding:10px;border-radius:8px}
small{color:#666}.disc{font-size:12px;color:#444}.row{display:grid;grid-template-columns:1fr 1fr;gap:10px}@media(max-width:560px){.row{grid-template-columns:1fr}}
</style>"""
def page(title,body,u=None):
    nav=('<a href="/app">My SWMS</a> · <a href="/logout">Log out</a>' if u else '<a href="/pro">PaintSWMS Pro</a>')
    ts='<script src="https://challenges.cloudflare.com/turnstile/v0/api.js" async defer></script>' if TS_SITE else ''
    return HTMLResponse(f"<!doctype html><html lang=en-AU><head><meta charset=utf-8><title>{esc(title)}</title>{CSS}{ts}<script src='https://gumroad.com/js/gumroad.js'></script>{TRACK_JS}</head><body><header><div class=w><a class=logo href='/'>Paint<span>SWMS</span></a><div>{nav}</div></div></header><div class=w>{body}</div></body></html>")
def ts_widget(): return f'<div class="cf-turnstile" data-sitekey="{TS_SITE}"></div>' if TS_SITE else ''
HP='<input class=hp name=website tabindex=-1 autocomplete=off>'
def seg(name,opts,val):
    return '<div class=seg>'+''.join(f'<label><input type=radio name={name} value={k} {"checked" if k==val else ""}><span>{v}</span></label>' for k,v in opts)+'</div>'
def wizard(action,d,btn,extra=""):
    st=''.join(f'<option {"selected" if s==d.get("state") else ""}>{s}</option>' for s in ["NSW","VIC","QLD","WA","SA","TAS","ACT","NT"])
    return f"""<form method=post action="{action}" class=card>{HP}{extra}
<label>Business name</label><input name=biz required maxlength=120 value="{esc(d.get('biz'))}" placeholder="e.g. Smith Painting">
<div class=row><div><label>Job site (address or suburb)</label><input name=site required maxlength=160 value="{esc(d.get('site'))}"></div><div><label>State</label><select name=state>{st}</select></div></div>
<label>Start date</label><input type=date name=start required value="{esc(d.get('start') or datetime.date.today().isoformat())}">
<label>Work at height</label>{seg('height',[('ladder','Ladder'),('scaffold','Scaffold'),('ewp','EWP'),('none','None')],d.get('height','ladder'))}
<label>Near overhead powerlines?</label>{seg('power',[('y','Yes'),('n','No')],d.get('power','n'))}
<label>Built before 1990? (asbestos / lead paint)</label>{seg('old',[('y','Yes / not sure'),('n','No')],d.get('old','y'))}
<label>Next to a road or traffic?</label>{seg('road',[('y','Yes'),('n','No')],d.get('road','n'))}
<p style="margin-top:18px"><button class=btn>{btn}</button></p></form>"""
FIELDS=["biz","site","state","start","height","power","old","road"]
def clean(f):
    d={k:str(f.get(k,""))[:160] for k in FIELDS}
    if d["state"] not in ["NSW","VIC","QLD","WA","SA","TAS","ACT","NT"]: d["state"]="NSW"
    if d["height"] not in C.HEIGHT: d["height"]="ladder"
    for k in ("power","old","road"): d[k]="y" if d[k]=="y" else "n"
    return d
LOCK_NAMES={"Work near overhead powerlines or energised services":"Powerlines near the eaves and gutters","Surface preparation on a pre-1990 building":"Pre-1990 prep: lead paint & asbestos","Work next to a road or traffic":"Traffic next to the work area"}
def locked_names(d):
    n=[LOCK_NAMES.get(t,t) for t,_,free in C.build(d)[1] if not free]
    if d.get("height")=="ewp": n.insert(0,"EWP rescue plan and exclusion zone")
    return n+["Site review and sign-off checklist"]
def doc_html(d,prof=None,signs=(),preview=False):
    tick,secs=C.build(d); prof=prof or {}
    dt=d.get("start","")
    try: dt=datetime.date.fromisoformat(dt).strftime("%-d %B %Y")
    except Exception: pass
    h=f"""<div class="card {'wm' if preview else ''}"><h2 style="margin:0">Safe Work Method Statement – Painting</h2><small>Structured to the Safe Work Australia model SWMS template.</small>
<table><tr><td>PCBU</td><td>{esc(d['biz'])}{(' · ABN '+esc(prof.get('abn'))) if prof.get('abn') else ''}</td><td>Principal contractor</td><td>__________</td></tr>
<tr><td>Workplace</td><td>{esc(d['site'])}, {esc(d['state'])}</td><td>Start date</td><td>{esc(dt)}</td></tr>
<tr><td>Contact</td><td>{esc(prof.get('contact') or '__________')} {esc(prof.get('phone') or '')}</td><td>Date given to PC</td><td>__________</td></tr></table>
<p style="font-size:13px">{C.NOTE}</p><h3>High risk construction work</h3><table>"""
    for i in range(0,18,3): h+="<tr>"+"".join(f"<td>{'☒' if i+j in tick else '☐'} {C.HRCW[i+j]}</td>" for j in range(3))+"</tr>"
    h+="</table>"
    for title,rows,free in secs:
        locked=preview and not free
        h+=f"<h3>{('🔒 '+esc(LOCK_NAMES.get(title,title))+' — locked') if locked else esc(title)}</h3><div class='{'lock' if locked else ''}'><table><tr><th>Task</th><th>Hazards and risks</th><th>Control measures</th><th>Who / monitored</th></tr>"
        for t,z,cm in rows: h+=f"<tr><td>{t}</td><td>{z}</td><td>{cm}</td><td>Supervisor: ______ · daily pre-start</td></tr>"
        h+="</table></div>"
        if locked: h+="<p><a class='btn sm' href='#unlock'>Unlock this job — $19</a></p>"
    h+=f"<h3>Compliance and review</h3><p style='font-size:13px'>{C.REVIEW}</p>"
    if prof.get("equipment"): h+=f"<p style='font-size:13px'><b>Usual equipment:</b> {esc(prof['equipment'])}</p>"
    h+="<h3>Worker consultation and sign-off</h3><table><tr><th>Name</th><th>Date</th><th>Signature</th></tr>"
    names=list(signs) or [(n.strip(),"") for n in (prof.get("crew") or "").split(",") if n.strip()]
    for n,dd in names: h+=f"<tr><td>{esc(n)}</td><td>{esc(dd)}</td><td></td></tr>"
    h+="<tr><td>&nbsp;</td><td></td><td></td></tr>"*max(1,4-len(names))
    h+=f"</table><p class=disc><b>{C.disclaimer(d['state'])}</b></p></div>"
    return h

def pdf_bytes(d,prof=None,signs=(),watermark=False):
    from reportlab.lib.pagesizes import A4,landscape
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    ss=getSampleStyleSheet(); P=lambda t,s="BodyText":Paragraph(html.escape(str(t)),ss[s]); prof=prof or {}
    tick,secs=C.build(d); el=[P("Safe Work Method Statement – Painting","Title")]
    def tb(rows,w=None):
        t=Table([[P(c) for c in r] for r in rows],colWidths=w,repeatRows=1); t.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.5,colors.grey),("VALIGN",(0,0),(-1,-1),"TOP")])); return t
    el+=[tb([["PCBU",d["biz"]+((" · ABN "+prof["abn"]) if prof.get("abn") else ""),"Principal contractor",""],["Workplace",f"{d['site']}, {d['state']}","Start date",d["start"]],["Contact",(prof.get("contact") or "")+" "+(prof.get("phone") or ""),"Date given to PC",""]]),Spacer(1,6),P(C.NOTE),P("High risk construction work","Heading3")]
    el.append(tb([[("[X] " if i+j in tick else "[ ] ")+C.HRCW[i+j] for j in range(3)] for i in range(0,18,3)]))
    for title,rows,free in secs:
        if watermark and not free: continue
        el+=[P(title,"Heading3"),tb([["Task","Hazards and risks","Control measures","Who / monitored"]]+[[a,b,c,"Supervisor: ______ · daily pre-start"] for a,b,c in rows],[120,170,350,120])]
    if watermark: el.append(P("The other sections for this job unlock for US$19 at "+BASE_HOST))
    el+=[P("Compliance and review","Heading3"),P(C.REVIEW),P("Worker consultation and sign-off","Heading3")]
    names=list(signs) or [(n.strip(),"") for n in (prof.get("crew") or "").split(",") if n.strip()]
    el.append(tb([["Name","Date","Signature"]]+[[n,x,""] for n,x in names]+[["","",""]]*max(1,4-len(names)),[250,150,300]))
    el+=[Spacer(1,8),P(C.disclaimer(d["state"]))]
    def wm(cv,doc):
        if watermark:
            cv.saveState(); cv.setFont("Helvetica-Bold",48); cv.setFillColorRGB(.91,.34,.16,alpha=.15); cv.translate(420,200); cv.rotate(25); cv.drawCentredString(0,0,"FREE PREVIEW - paintswms"); cv.restoreState(); cv.setFont("Helvetica-Bold",9); cv.drawString(28,12,"Free preview - paintswms. Unlock the full, clean SWMS for this job: "+BASE_HOST)
    b=io.BytesIO(); SimpleDocTemplate(b,pagesize=landscape(A4),leftMargin=28,rightMargin=28,topMargin=24,bottomMargin=24).build(el,onFirstPage=wm,onLaterPages=wm); return b.getvalue()
def docx_bytes(d,prof=None,signs=()):
    from docx import Document
    tick,secs=C.build(d); prof=prof or {}; doc=Document(); doc.add_heading("Safe Work Method Statement – Painting",0)
    t=doc.add_table(rows=0,cols=4); t.style="Table Grid"
    for r in [["PCBU",d["biz"]+((" · ABN "+prof["abn"]) if prof.get("abn") else ""),"Principal contractor",""],["Workplace",f"{d['site']}, {d['state']}","Start date",d["start"]],["Contact",(prof.get("contact") or "")+" "+(prof.get("phone") or ""),"Date given to PC",""]]:
        c=t.add_row().cells; [setattr(c[i],"text",r[i]) for i in range(4)]
    doc.add_paragraph(C.NOTE); doc.add_heading("High risk construction work",2)
    t=doc.add_table(rows=0,cols=3); t.style="Table Grid"
    for i in range(0,18,3):
        c=t.add_row().cells
        for j in range(3): c[j].text=("☒ " if i+j in tick else "☐ ")+C.HRCW[i+j]
    for title,rows,_ in secs:
        doc.add_heading(title,2); t=doc.add_table(rows=1,cols=4); t.style="Table Grid"
        for c,x in zip(t.rows[0].cells,["Task","Hazards and risks","Control measures","Who / monitored"]): c.text=x
        for a,b,cm in rows:
            c=t.add_row().cells; c[0].text=a; c[1].text=b; c[2].text=cm; c[3].text="Supervisor: ______ · daily pre-start"
    doc.add_heading("Compliance and review",2); doc.add_paragraph(C.REVIEW); doc.add_heading("Worker consultation and sign-off",2)
    t=doc.add_table(rows=1,cols=3); t.style="Table Grid"
    for c,x in zip(t.rows[0].cells,["Name","Date","Signature"]): c.text=x
    names=list(signs) or [(n.strip(),"") for n in (prof.get("crew") or "").split(",") if n.strip()]
    for n,x in names: c=t.add_row().cells; c[0].text=n; c[1].text=x
    for _ in range(max(1,4-len(names))): t.add_row()
    doc.add_paragraph(C.disclaimer(d["state"])); b=io.BytesIO(); doc.save(b); return b.getvalue()

# ---------- Layer A: free generator ----------
@app.middleware("http")
async def did(req,call_next):
    r=await call_next(req)
    if not req.cookies.get("did"): r.set_cookie("did",secrets.token_urlsafe(16),max_age=365*86400,httponly=True,samesite="lax")
    return r
@app.post("/t/{ev}")
def t_event(ev:str):
    if ev in ("checkout_click_pack","checkout_click_unlimited"): track(ev)
    return Response(status_code=204)
@app.get("/internal/funnel")
@app.get("/metrics")
def metrics(req:Request):
    tok=os.environ.get("METRICS_TOKEN","")
    if not tok: return Response(status_code=404)
    if not hmac.compare_digest(req.headers.get("x-metrics-token",""),tok): return Response(status_code=401)
    with db() as c: rows=c.execute("select day,event,n from counters order by day").fetchall()
    days={}; tot={e:0 for e in EVENTS}
    for r in rows:
        days.setdefault(r["day"],{e:0 for e in EVENTS})[r["event"]]=r["n"]; tot[r["event"]]=tot.get(r["event"],0)+r["n"]
    return {"days":days,"all_time":tot}
@app.get("/health")
def health(): return {"ok":True}
@app.get("/",response_class=HTMLResponse)
def home(req:Request):
    track("visit_home")
    return page("PaintSWMS – painting SWMS in 2 minutes",f"""<h1 style="margin:18px 0 4px">Answer 6 questions, get your painting SWMS</h1>
<p>No account. Your SWMS comes up on screen straight away, with your business and job filled in. Built on the Safe Work Australia model template.</p>
{wizard('/preview',{},'Show my SWMS')}
<div class=card><h3 style="margin-top:0">Need one for every job?</h3><p>PaintSWMS Pro saves your business details once, keeps a library of every SWMS, tracks worker sign-offs, and gives you clean PDF and editable Word files. <a href="/pro">See Pro</a></p></div>
<p class=disc>PaintSWMS is a template tool, not legal or safety advice. Tailor every SWMS to the site, consult your workers and check your state or territory regulator's requirements. <a href="/privacy">Privacy</a></p>
<!-- PLACEHOLDER (wording not approved by owner, do not enable): <p>If a principal contractor rejects your PaintSWMS SWMS, send us their reason within 14 days and we'll refund you.</p> -->""",user(req))
@app.post("/preview",response_class=HTMLResponse)
async def preview(req:Request):
    f=await req.form()
    if f.get("website"): return page("Thanks","<p class=ok>Thanks.</p>")
    if limited([ip(req),device(req)],"gen",20,3600): return page("Slow down","<p class=err>Too many previews from this connection. Try again in an hour.</p>")
    d=clean(f); tok=sign({"d":d,"t":now()}); track("preview_generated")
    locks="".join(f"<li>🔒 {esc(n)} — locked</li>" for n in locked_names(d))
    return page("Your SWMS preview",f"""<div class="card ok"><b>Here's your SWMS for {esc(d['site'])}.</b> Work at height and the sign-off sheet are free.</div>
<div class=card id=unlock><h2 style="margin-top:0">Unlock your full SWMS for this job — $19</h2><p>Still locked for this job:</p><ul>{locks}</ul>
<form method=post action="/unlock/start"><input type=hidden name=tok value='{esc(tok)}'><button class=btn>Unlock your full SWMS for this job — $19</button></form>
<small>Your answers are already filled in. You get a clean PDF and an editable Word file straight away, saved to your account for re-downloads.</small></div>
<details class=card><summary><b>Just want the free part? Download the watermarked preview</b></summary>
<form method=post action="/download/start" class=card>{HP}<input type=hidden name=tok value='{esc(tok)}'><label>Email</label><input type=email name=email required maxlength=120>{ts_widget()}<p><button class=btn>Email me a download code</button></p><small>We'll email you a 6-digit code. We only use your email to send your SWMS and the odd PaintSWMS update.</small></form></details>
{doc_html(d,preview=True)}""")
@app.post("/download/start",response_class=HTMLResponse)
async def dl_start(req:Request):
    f=await req.form(); email=str(f.get("email","")).strip().lower(); d=(unsign(str(f.get("tok","")))or{}).get("d")
    if f.get("website") or not d: return page("Thanks","<p class=ok>Thanks.</p>")
    if "@" not in email or is_disposable(email): return page("Use a real email","<p class=err>Please use your normal work or personal email. Temporary inboxes aren't accepted.</p>")
    if not turnstile_ok(f.get("cf-turnstile-response",""),req): return page("Check failed","<p class=err>The human check didn't pass. Go back and try again.</p>")
    if limited([ip(req),device(req)],"code",5,3600) or limited([email],"code_e",3,3600): return page("Slow down","<p class=err>Too many code requests. Try again in an hour.</p>")
    with db() as c: c.execute("insert into leads values(?,?,?,?)",(email,now(),ip(req),json.dumps(d)))
    track("email_captured")
    if not email_ready() and not DEV:
        tok=sign({"d":d,"e":email,"t":now(),"g":1})
        return page("Download",f"""<div class=card><p>Grab the free download through our Gumroad page (it's free, and it's how we keep our mailing list), then come back and download your filled-in PDF.</p>
<p><a class=btn href="{FREE_URL}?wanted=true&email={esc(email)}" target=_blank>Open the free Gumroad page</a></p><p><a class="btn sec" href="/download/pdf?tok={esc(tok)}">Download my PDF</a></p></div>""")
    code=make_code(email,"dl",json.dumps(d)); send_code(email,code)
    dev=f"<p class=ok>DEV MODE code: {code}</p>" if DEV else ""
    return page("Check your email",f"""{dev}<form method=post action="/download/verify" class=card><input type=hidden name=email value="{esc(email)}"><label>Enter the 6-digit code we sent to {esc(email)}</label><input name=code inputmode=numeric autocomplete=one-time-code maxlength=6 required><p><button class=btn>Verify and download</button></p></form>""")
@app.post("/download/verify",response_class=HTMLResponse)
async def dl_verify(req:Request):
    f=await req.form(); email=str(f.get("email","")).lower(); p=check_code(email,"dl",str(f.get("code","")))
    if p is None: return page("Code didn't work","<p class=err>That code is wrong or expired. Go back and request a new one.</p>")
    d=json.loads(p); tok=sign({"d":d,"e":email,"t":now(),"v":1})
    return page("Your download",f"""<div class="card ok"><b>Verified.</b></div><div class=card><p><a class=btn href="/download/pdf?tok={esc(tok)}">Download my free PDF</a></p></div>
<div class=card><h2 style="margin-top:0">Want every section for this job?</h2><p>Your answers carry over. Unlock the full SWMS for this job and you get a clean PDF and Word file, saved to your account with a sign-off list for your crew.</p><form method=post action="/unlock/start"><input type=hidden name=tok value='{esc(tok)}'><button class=btn>Unlock your full SWMS for this job — $19</button></form><p><small>Or make one for every job: <a href="/pro?tok={esc(tok)}">Unlimited, $9/month</a>.</small></p></div>""")
@app.get("/download/pdf")
def dl_pdf(req:Request,tok:str=""):
    t=unsign(tok)
    if not t or not (t.get("v") or t.get("g")) or now()-t["t"]>3600: return page("Link expired","<p class=err>This download link has expired. Make the SWMS again from the home page.</p>")
    if limited([ip(req)],"pdf",15,3600): return page("Slow down","<p class=err>Too many downloads. Try again in an hour.</p>")
    return Response(pdf_bytes(t["d"],watermark=True),media_type="application/pdf",headers={"Content-Disposition":"attachment; filename=PaintSWMS_free_preview.pdf"})

# ---------- Layer B: Pro ----------
@app.get("/pro",response_class=HTMLResponse)
def pro(req:Request,tok:str=""):
    u=user(req)
    if u: return RedirectResponse("/app"+(f"?tok={tok}" if tok else ""),303)
    t=unsign(tok) or {}; em=t.get("e","")
    login=(f"""<form method=post action="/login/start" class=card>{HP}<input type=hidden name=tok value='{esc(tok)}'><h3 style="margin-top:0">Log in or start</h3><label>Email</label><input type=email name=email required value="{esc(em)}">{ts_widget()}<p><button class=btn>Email me a login code</button></p><small>No password. We email you a 6-digit code.</small></form>""" if (email_ready() or DEV) else "")
    return page("PaintSWMS Pro",f"""<h1>PaintSWMS Pro</h1><div class=card><ul><li>Business profile saved once: name, ABN, contact, crew, usual gear</li><li>Unlimited site-specific SWMS from the same 6 questions</li><li>Library of every job: view, edit, duplicate for the next one</li><li>Worker sign-off list on each SWMS</li><li>Clean PDF and editable Word, every relevant section included</li></ul>
<p><a class="btn gumroad-button" href="{PRO_URL}">Get Unlimited — US$9/month</a></p><p><small>Only need one job? <a href="/">Answer 6 questions and unlock that SWMS for US$19</a>. Checkout by Gumroad.</small></p></div>
{login}<form method=post action="/login/license" class=card>{HP}<input type=hidden name=tok value='{esc(tok)}'><h3 style="margin-top:0">Already bought? Log in with your licence key</h3><label>Gumroad licence key</label><input name=key required maxlength=64 placeholder="XXXXXXXX-XXXXXXXX-XXXXXXXX-XXXXXXXX"><p><button class="btn sec">Unlock Pro</button></p><small>The key is in your Gumroad receipt. Your account uses the email you bought with.</small></form>""")
@app.post("/login/start",response_class=HTMLResponse)
async def login_start(req:Request):
    f=await req.form(); email=str(f.get("email","")).strip().lower()
    if f.get("website"): return page("Thanks","<p class=ok>Thanks.</p>")
    if "@" not in email or is_disposable(email): return page("Use a real email","<p class=err>Temporary inboxes aren't accepted.</p>")
    if not turnstile_ok(f.get("cf-turnstile-response",""),req): return page("Check failed","<p class=err>Human check failed.</p>")
    if limited([ip(req),device(req)],"login",8,3600) or limited([email],"login_e",4,3600): return page("Slow down","<p class=err>Too many attempts. Try again in an hour.</p>")
    code=make_code(email,"login",str(f.get("tok",""))); ok=send_code(email,code)
    dev=f"<p class=ok>DEV MODE code: {code}</p>" if DEV else ""
    return page("Check your email",f"""{dev}<form method=post action="/login/verify" class=card><input type=hidden name=email value="{esc(email)}"><label>6-digit code sent to {esc(email)}</label><input name=code inputmode=numeric autocomplete=one-time-code maxlength=6 required><p><button class=btn>Log in</button></p></form>""")
@app.post("/login/verify")
async def login_verify(req:Request):
    f=await req.form(); email=str(f.get("email","")).lower(); p=check_code(email,"login",str(f.get("code","")))
    if p is None: return page("Code didn't work","<p class=err>Wrong or expired code.</p>")
    return login_resp(email,"/app"+(f"?tok={p}" if p and p!="{}" else ""))
def gumroad_sale(sale_id):
    tokn=os.environ.get("GUMROAD_ACCESS_TOKEN")
    if not tokn or not sale_id: return None
    try: j=httpx.get(f"https://api.gumroad.com/v2/sales/{sale_id}",params={"access_token":tokn},timeout=10).json()
    except Exception: return None
    s=j.get("sale") if j.get("success") else None
    if not s or s.get("refunded") or s.get("chargedback") or s.get("product_id") not in GUMROAD_PRODUCTS: return None
    return s.get("product_id"),s
def claim(pid,p,key,d):
    """Turn a verified purchase into an account + saved SWMS. Returns (response) ."""
    email=(p.get("email") or "").lower()
    with db() as c:
        used=c.execute("select * from unlocks where key=?",(key,)).fetchone()
    if pid==PACK_ID and used:
        r=login_resp(email,f"/done/{used['swms_id']}"); return r
    r0=login_resp(email,"/")
    with db() as c:
        uid=c.execute("select id from users where email=?",(email,)).fetchone()["id"]
        sid=None
        if d: sid=c.execute("insert into swms(user_id,data,created,updated,paid) values(?,?,?,?,1) returning id",(uid,json.dumps(d),now(),now())).fetchone()["id"]
        if pid==PACK_ID:
            if sid: c.execute("insert into unlocks values(?,?,?,?) on conflict do nothing",(key,sid,uid,now()))
        else: c.execute("update users set pro_until=?,license=? where id=?",(now()+35*86400,key,uid))
    track("unlock_success"); r0.headers["location"]=f"/done/{sid}" if sid else "/app"; r0.delete_cookie("pend"); return r0
def pend_data(pid_):
    with db() as c: r=c.execute("select data from pending where id=?",(pid_ or "",)).fetchone()
    return json.loads(r["data"]) if r else None
@app.post("/unlock/start")
async def unlock_start(req:Request):
    f=await req.form(); d=(unsign(str(f.get("tok","")))or{}).get("d")
    if not d: return RedirectResponse("/",303)
    if limited([ip(req),device(req)],"pend",20,3600): return page("Slow down","<p class=err>Too many attempts. Try again in an hour.</p>")
    pid_=secrets.token_urlsafe(12)
    with db() as c: c.execute("insert into pending values(?,?,?)",(pid_,json.dumps(d),now()))
    r=RedirectResponse(f"/unlock/{pid_}",303); r.set_cookie("pend",pid_,max_age=7*86400,httponly=True,samesite="lax"); return r
@app.get("/unlock/return")
def unlock_return(req:Request,sale_id:str=""):
    v=gumroad_sale(sale_id); d=pend_data(req.cookies.get("pend"))
    if v and v[0]!=PACK_ID: return claim(v[0],v[1],"sale:"+sale_id,d)
    if v and not d: return page("Open on the same device","<div class=card><p>Payment confirmed. Open this button on the phone or computer where you filled in your SWMS, or go to <a href='/'>PaintSWMS</a>, fill in the job again and paste the licence key from your Gumroad receipt on the unlock page.</p></div>")
    if not v or not d: return RedirectResponse(f"/unlock/{req.cookies.get('pend','')}" if d else "/",303)
    return claim(v[0],v[1],"sale:"+sale_id,d)
@app.get("/unlock/{pid_}",response_class=HTMLResponse)
def unlock_page(req:Request,pid_:str):
    d=pend_data(pid_)
    if not d: return RedirectResponse("/",303)
    locks="".join(f"<li>{esc(n)}</li>" for n in locked_names(d))
    return page("Unlock your SWMS",f"""<div class=card><h2 style="margin-top:0">Unlock your full SWMS for this job — $19</h2>
<p><b>{esc(d['biz'])}</b> · {esc(d['site'])}, {esc(d['state'])} · starts {esc(d['start'])}</p><p>Adds every section for this job:</p><ul>{locks}</ul>
<p><b>1.</b> Pay with Gumroad (card, Apple Pay, Google Pay, PayPal).</p><p><a class="btn gumroad-button" href="{PACK_URL}?wanted=true" data-gumroad-overlay-checkout="true">Pay $19 and unlock</a></p>
<p><b>2.</b> Paste the licence key from your Gumroad receipt. Your SWMS downloads straight away and is saved to an account under your purchase email. No password.</p>
<form method=post action="/unlock/{esc(pid_)}">{HP}<input name=key required maxlength=64 placeholder="XXXXXXXX-XXXXXXXX-XXXXXXXX-XXXXXXXX" autocomplete=off><p><button class="btn sec">Get my SWMS</button></p></form>
<small>One $19 unlock covers one SWMS for one job. Already on Unlimited? Paste that key instead.</small></div>
<p class=disc>{C.disclaimer(d['state'])}</p>""",user(req))
@app.post("/unlock/{pid_}")
async def unlock_claim(req:Request,pid_:str):
    f=await req.form(); d=pend_data(pid_)
    if f.get("website") or not d: return RedirectResponse("/",303)
    if limited([ip(req),device(req)],"lic",10,3600): return page("Slow down","<p class=err>Too many attempts. Try again in an hour.</p>")
    key=str(f.get("key","")).strip(); v=verify_license(key)
    if not v: return page("Key not recognised",f"<p class=err>That licence key isn't valid, or the purchase was refunded. <a href='/unlock/{esc(pid_)}'>Try again</a></p>")
    with db() as c: c.execute("delete from pending where id=?",(pid_,))
    return claim(v[0],v[1],key,d)
UPSELL=f"""<div class=card><h3 style="margin-top:0">Need one for every job? Unlimited — $9/month</h3><p>Every job, every section, clean PDF and Word, your business details filled in each time. Cancel any time in Gumroad.</p><p><a class="btn gumroad-button" href="{PRO_URL}">Go Unlimited — $9/month</a></p><small>After paying, paste your Unlimited key in <a href="/app">My SWMS</a>.</small></div>"""
def can(u,r): return bool(r and (r["paid"] or is_pro(u)))
@app.get("/done/{sid}",response_class=HTMLResponse)
def done(req:Request,sid:int):
    u,red=need(req)
    if red: return red
    r=own(u,sid)
    if not can(u,r): return RedirectResponse("/app",303)
    d=json.loads(r["data"])
    return page("Your SWMS is ready",f"""<div class="card ok"><h2 style="margin-top:0">Your SWMS is ready</h2><p>{esc(d['site'])}, {esc(d['state'])}. Saved to your account ({esc(u['email'])}) so you can download it again any time.</p></div>
<div class=card><p><a class=btn href="/swms/{sid}/pdf">Download PDF</a></p><p><a class="btn sec" href="/swms/{sid}/docx">Download Word (editable)</a></p><p><a href="/swms/{sid}">Open it to add worker sign-offs</a></p></div>{"" if is_pro(u) else UPSELL}""",u)
@app.post("/login/license")
async def login_license(req:Request):
    f=await req.form()
    if f.get("website"): return page("Thanks","<p class=ok>Thanks.</p>")
    if limited([ip(req),device(req)],"lic",10,3600): return page("Slow down","<p class=err>Too many attempts.</p>")
    key=str(f.get("key","")).strip(); v=verify_license(key)
    if not v: return page("Key not recognised","<p class=err>That licence key isn't valid, or the purchase was refunded or the subscription has ended.</p>")
    return claim(v[0],v[1],key,(unsign(str(f.get("tok","")))or{}).get("d"))
@app.post("/unlock")
async def unlock_pro(req:Request):
    u=user(req)
    if not u: return RedirectResponse("/pro",303)
    f=await req.form(); key=str(f.get("key","")).strip(); v=verify_license(key)
    if not v or v[0]==PACK_ID: return page("Key not recognised","<p class=err>That isn't an active Unlimited key. (A $19 job key unlocks the SWMS it was bought for.)</p>",u)
    with db() as c: c.execute("update users set pro_until=?,license=? where id=?",(now()+35*86400,key,u["id"]))
    track("unlock_success"); return RedirectResponse("/app",303)
@app.get("/logout")
def logout():
    r=RedirectResponse("/",303); r.delete_cookie("sess"); return r
def prof_of(uid):
    with db() as c: r=c.execute("select * from profiles where user_id=?",(uid,)).fetchone()
    return dict(r) if r else {}
def need(req):
    u=user(req)
    if not u: return None,RedirectResponse("/pro",303)
    if u["license"] and u["pro_until"]<now()+86400:
        v=verify_license(u["license"])
        if v and v[0]!=PACK_ID:
            with db() as c: c.execute("update users set pro_until=? where id=?",(now()+35*86400,u["id"]))
            u=user(req)
    return u,None
def own(u,sid):
    with db() as c: return c.execute("select * from swms where id=? and user_id=?",(sid,u["id"])).fetchone()
@app.get("/app",response_class=HTMLResponse)
def dash(req:Request,tok:str=""):
    u,red=need(req)
    if red: return red
    t=unsign(tok) or {}
    if t.get("d") and is_pro(u):
        with db() as c: c.execute("insert into swms(user_id,data,created,updated) values(?,?,?,?)",(u["id"],json.dumps(t["d"]),now(),now()))
        return RedirectResponse("/app",303)
    pro_=is_pro(u); p=prof_of(u["id"])
    with db() as c: rows=c.execute("select * from swms where user_id=? order by updated desc",(u["id"],)).fetchall()
    def act(r):
        if not can(u,r): return "<small>Preview only</small>"
        a=f"<a href='/swms/{r['id']}'>Open</a> · <a href='/swms/{r['id']}/pdf'>PDF</a> · <a href='/swms/{r['id']}/docx'>Word</a>"
        if pro_: a+=f" · <a href='/swms/{r['id']}/edit'>Edit</a> · <form method=post action='/swms/{r['id']}/dup' style='display:inline'><button class='btn sm sec'>Duplicate</button></form>"
        return a
    lst="".join(f"<tr><td>{esc(json.loads(r['data'])['site'])}</td><td>{esc(json.loads(r['data'])['start'])}</td><td>{act(r)}</td></tr>" for r in rows) or "<tr><td colspan=3>No SWMS yet.</td></tr>"
    if pro_: top=f"<p><a class=btn href='/swms/new'>+ New SWMS</a></p>"+("" if p.get("business") else "<p class=ok>Tip: <a href='/profile'>add your business details</a> once and every new SWMS fills them in.</p>")
    else: top=f"""<p><a class=btn href="/">+ New SWMS for another job ($19)</a></p>{UPSELL}<form method=post action="/unlock" class=card><label>Bought Unlimited? Paste the key</label><input name=key required maxlength=64><p><button class="btn sm sec">Activate Unlimited</button></p></form>"""
    return page("My SWMS",f"""{top}<div class=card><h3 style="margin-top:0">My SWMS</h3><table><tr><th>Job site</th><th>Start</th><th></th></tr>{lst}</table></div><p><a href="/profile">Business profile</a></p>""",u)
@app.get("/profile",response_class=HTMLResponse)
def profile(req:Request):
    u,red=need(req)
    if red: return red
    p=prof_of(u["id"]); fld=lambda k,l,ph="":f"<label>{l}</label><input name={k} value='{esc(p.get(k))}' placeholder='{ph}' maxlength=300>"
    return page("Business profile",f"<form method=post class=card>{fld('business','Business name')}{fld('abn','ABN (optional)')}{fld('contact','Contact person')}{fld('phone','Phone')}{fld('crew','Default crew (comma separated)','Dave Smith, Priya N')}{fld('equipment','Usual equipment','Extension ladders, 4 m mobile scaffold, 12 m boom lift')}<label>Usual state</label><input name=state value='{esc(p.get('state') or 'NSW')}' maxlength=3><p><button class=btn>Save</button></p></form>",u)
@app.post("/profile")
async def profile_save(req:Request):
    u,red=need(req)
    if red: return red
    f=await req.form(); v=[str(f.get(k,""))[:300] for k in ["business","abn","contact","phone","crew","equipment","state"]]
    with db() as c: c.execute("insert into profiles values(?,?,?,?,?,?,?,?) on conflict(user_id) do update set business=excluded.business,abn=excluded.abn,contact=excluded.contact,phone=excluded.phone,crew=excluded.crew,equipment=excluded.equipment,state=excluded.state",(u["id"],*v))
    return RedirectResponse("/app",303)
@app.get("/swms/new",response_class=HTMLResponse)
def new(req:Request):
    u,red=need(req)
    if red: return red
    if not is_pro(u): return RedirectResponse("/",303)
    p=prof_of(u["id"]); return page("New SWMS",wizard("/swms/new",{"biz":p.get("business",""),"state":p.get("state","NSW")},"Create SWMS"),u)
@app.post("/swms/new")
async def new_save(req:Request):
    u,red=need(req)
    if red or not is_pro(u): return red or RedirectResponse("/",303)
    d=clean(await req.form())
    with db() as c: nid=c.execute("insert into swms(user_id,data,created,updated) values(?,?,?,?) returning id",(u["id"],json.dumps(d),now(),now())).fetchone()["id"]
    return RedirectResponse(f"/swms/{nid}",303)
def signs_of(sid):
    with db() as c: return [(r["name"],r["dt"]) for r in c.execute("select * from signoffs where swms_id=? order by id",(sid,)).fetchall()]
@app.get("/swms/{sid}",response_class=HTMLResponse)
def view(req:Request,sid:int):
    u,red=need(req)
    if red: return red
    r=own(u,sid)
    if not can(u,r): return RedirectResponse("/app",303)
    d=json.loads(r["data"]); ed=f' <a href="/swms/{sid}/edit">Edit answers</a>' if is_pro(u) else ""
    return page("SWMS",f"""<p><a class="btn sm" href="/swms/{sid}/pdf">PDF</a> <a class="btn sm sec" href="/swms/{sid}/docx">Word</a>{ed}</p>
<form method=post action="/swms/{sid}/sign" class=card><h3 style="margin-top:0">Add worker sign-off</h3><div class=row><input name=name required maxlength=80 placeholder="Worker name"><input type=date name=dt value="{datetime.date.today().isoformat()}"></div><p><button class="btn sm">Add</button></p></form>
{doc_html(d,prof_of(u['id']),signs_of(sid))}""",u)
@app.get("/swms/{sid}/edit",response_class=HTMLResponse)
def edit(req:Request,sid:int):
    u,red=need(req)
    if red: return red
    r=own(u,sid)
    if not r or not is_pro(u): return RedirectResponse("/app",303)
    return page("Edit SWMS",wizard(f"/swms/{sid}/edit",json.loads(r["data"]),"Save changes"),u)
@app.post("/swms/{sid}/edit")
async def edit_save(req:Request,sid:int):
    u,red=need(req)
    if red or not own(u,sid) or not is_pro(u): return RedirectResponse("/app",303)
    d=clean(await req.form())
    with db() as c: c.execute("update swms set data=?,updated=? where id=?",(json.dumps(d),now(),sid))
    return RedirectResponse(f"/swms/{sid}",303)
@app.post("/swms/{sid}/dup")
def dup(req:Request,sid:int):
    u,red=need(req)
    r=None if red else own(u,sid)
    if not r or not is_pro(u): return RedirectResponse("/app",303)
    d=json.loads(r["data"]); d["site"]=""; d["start"]=datetime.date.today().isoformat()
    return page("Duplicate for a new job",wizard("/swms/new",d,"Create SWMS for new job"),u)
@app.post("/swms/{sid}/sign")
async def sign_add(req:Request,sid:int):
    u,red=need(req)
    if red or not can(u,own(u,sid)): return RedirectResponse("/app",303)
    f=await req.form()
    with db() as c: c.execute("insert into signoffs(swms_id,name,dt) values(?,?,?)",(sid,str(f.get("name",""))[:80],str(f.get("dt",""))[:10]))
    return RedirectResponse(f"/swms/{sid}",303)
@app.get("/swms/{sid}/{fmt}")
def export(req:Request,sid:int,fmt:str):
    u,red=need(req)
    r=None if red else own(u,sid)
    if not can(u,r): return RedirectResponse("/app",303)
    d=json.loads(r["data"]); p=prof_of(u["id"]); s=signs_of(sid); name="SWMS_"+"".join(ch for ch in d["site"] if ch.isalnum())[:30]
    if fmt=="pdf": return Response(pdf_bytes(d,p,s),media_type="application/pdf",headers={"Content-Disposition":f"attachment; filename={name}.pdf"})
    if fmt=="docx": return Response(docx_bytes(d,p,s),media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",headers={"Content-Disposition":f"attachment; filename={name}.docx"})
    return RedirectResponse("/app",303)
@app.get("/privacy",response_class=HTMLResponse)
def privacy(req:Request):
    return page("Privacy","<div class=card><h2>Privacy</h2><p>We collect your email when you download a free SWMS or log in, the answers you enter, your business profile, and the names you add to sign-off lists. We keep your IP address and a device cookie for a short time to stop abuse. Payments are handled by Gumroad. We don't sell your details. Email us to have your account deleted.</p></div>")
